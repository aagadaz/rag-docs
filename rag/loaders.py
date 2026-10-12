import csv
import json
import re
import unicodedata
from pathlib import Path
from typing import Callable

import requests
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from langchain_core.documents import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

LoaderFn = Callable[[Path], list[Document]]
_REGISTRY: dict[str, LoaderFn] = {}

ROWS_PER_DOCUMENT = 50
CODE_EXTENSIONS = {
    ".py": "python", ".js": "javascript", ".ts": "typescript", ".tsx": "tsx",
    ".jsx": "jsx", ".java": "java", ".go": "go", ".rs": "rust", ".c": "c",
    ".cpp": "cpp", ".sql": "sql", ".sh": "bash", ".yaml": "yaml",
    ".yml": "yaml", ".toml": "toml",
}


class UnsupportedFormatError(ValueError):
    pass


def register(*extensions: str):
    def decorator(fn: LoaderFn) -> LoaderFn:
        for ext in extensions:
            _REGISTRY[ext.lower()] = fn
        return fn
    return decorator


def supported_extensions() -> list[str]:
    return sorted(_REGISTRY)


# ---------- outils communs ----------

def _clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"-\n(\w)", r"\1", text)
    text = re.sub(r"(?<=\S)[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _read_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="latin-1")


def _single_document(text: str, source: str, file_type: str, **extra) -> list[Document]:
    text = _clean_text(text)
    if not text:
        raise ValueError(f"Aucun texte extractible dans {source}.")
    return [Document(
        page_content=text,
        metadata={"source": source, "file_type": file_type, **extra},
    )]


def _rows_to_documents(header, rows, source, file_type, **extra) -> list[Document]:
    header = [str(h).strip() if h is not None else f"col{i + 1}" for i, h in enumerate(header)]
    lines = []
    for row in rows:
        pairs = [
            f"{name}: {value}"
            for name, value in zip(header, row)
            if value is not None and str(value).strip()
        ]
        if pairs:
            lines.append("; ".join(pairs))

    documents = []
    for start in range(0, len(lines), ROWS_PER_DOCUMENT):
        batch = lines[start:start + ROWS_PER_DOCUMENT]
        documents.append(Document(
            page_content="\n".join(batch),
            metadata={
                "source": source,
                "file_type": file_type,
                "row_start": start + 1,
                "row_end": start + len(batch),
                **extra,
            },
        ))
    return documents


# ---------- texte et code ----------

@register(".txt", ".md", ".rst")
def _load_text(path: Path) -> list[Document]:
    return _single_document(_read_text_file(path), path.name, path.suffix.lower().lstrip("."))


@register(*CODE_EXTENSIONS)
def _load_code(path: Path) -> list[Document]:
    text = _read_text_file(path).strip()
    if not text:
        raise ValueError(f"Fichier vide : {path.name}")
    return [Document(
        page_content=text,
        metadata={
            "source": path.name,
            "file_type": "code",
            "language": CODE_EXTENSIONS[path.suffix.lower()],
        },
    )]


@register(".json")
def _load_json(path: Path) -> list[Document]:
    raw = _read_text_file(path)
    try:
        text = json.dumps(json.loads(raw), indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        text = raw
    return [Document(
        page_content=text.strip(),
        metadata={"source": path.name, "file_type": "json"},
    )]


# ---------- bureautique ----------

@register(".pdf")
def _load_pdf(path: Path) -> list[Document]:
    reader = PdfReader(path)
    documents = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = _clean_text(page.extract_text() or "")

        if not text:
            continue

        documents.append(Document(
            page_content=text,
            metadata={"source": path.name, "file_type": "pdf", "page": page_number},
        ))

    if not documents:
        raise ValueError(
            f"Aucun texte extractible dans {path.name}. "
            "C'est probablement un PDF scanné (images), l'OCR n'est pas supporté."
        )

    return documents


@register(".docx")
def _load_docx(path: Path) -> list[Document]:
    docx_file = DocxDocument(path)
    parts = []

    for paragraph in docx_file.paragraphs:
        if not paragraph.text.strip():
            continue

        if paragraph.style.name.startswith("Heading"):
            line = "# " + paragraph.text
        else:
            line = paragraph.text

        parts.append(line)

    for table in docx_file.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cells.append(cell.text.strip())
            parts.append(" | ".join(cells))

    return _single_document("\n".join(parts), path.name, "docx")


@register(".pptx")
def _load_pptx(path: Path) -> list[Document]:
    presentation = Presentation(path)
    documents = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts = []

        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Notes : " + notes)

        text = _clean_text("\n".join(parts))
        if not text:
            continue

        documents.append(Document(
            page_content=text,
            metadata={"source": path.name, "file_type": "pptx", "slide": slide_number},
        ))

    if not documents:
        raise ValueError(f"Aucun texte extractible dans {path.name}.")

    return documents


# ---------- données tabulaires ----------

@register(".csv")
def _load_csv(path: Path) -> list[Document]:
    rows = list(csv.reader(_read_text_file(path).splitlines()))
    if len(rows) < 2:
        raise ValueError(f"CSV vide ou sans données : {path.name}")
    return _rows_to_documents(rows[0], rows[1:], path.name, "csv")


@register(".xlsx")
def _load_xlsx(path: Path) -> list[Document]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    documents = []

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if len(rows) < 2:
            continue
        documents.extend(
            _rows_to_documents(rows[0], rows[1:], path.name, "xlsx", sheet=sheet.title)
        )

    workbook.close()

    if not documents:
        raise ValueError(f"Aucune donnée dans {path.name}.")

    return documents


# ---------- web ----------

BLOCK_TAGS = ["p", "div", "section", "li", "tr", "br", "h1", "h2", "h3",
              "h4", "h5", "h6", "table", "ul", "ol", "dt", "dd", "blockquote"]
NOISE_TAGS = ["script", "style", "nav", "footer", "header", "aside", "noscript", "form"]
NOISE_SELECTORS = [".headerlink", ".related", ".sphinxsidebar", ".footer", "[role=navigation]"]


def _html_to_text(html: str) -> tuple[str, str]:
    soup = BeautifulSoup(html, "html.parser")
    title = soup.title.get_text(strip=True) if soup.title else ""

    for tag in soup(NOISE_TAGS):
        tag.decompose()
    for tag in soup.select(", ".join(NOISE_SELECTORS)):
        tag.decompose()

    main = (soup.find("main") or soup.find(attrs={"role": "main"})
            or soup.find("article") or soup.body or soup)

    for pre in main.find_all("pre"):
        pre.replace_with("\n```\n" + pre.get_text() + "\n```\n")

    for cell in main.find_all(["td", "th"]):
        cell.append(" | ")

    for tag in main.find_all(BLOCK_TAGS):
        tag.append("\n")

    return title, main.get_text()

@register(".html", ".htm")
def _load_html(path: Path) -> list[Document]:
    title, text = _html_to_text(_read_text_file(path))
    return _single_document(text, path.name, "html", title=title or path.stem)


def _load_url(url: str) -> list[Document]:
    response = requests.get(
        url,
        timeout=30,
        headers={"User-Agent": "rag-doc-assistant/1.0"},
    )
    response.raise_for_status()
    title, text = _html_to_text(response.text)
    return _single_document(text, url, "url", title=title or url)


# ---------- point d'entrée ----------

def load_source(source: str | Path) -> list[Document]:
    if isinstance(source, str) and source.startswith(("http://", "https://")):
        return _load_url(source)

    path = Path(source)
    if not path.exists():
        raise FileNotFoundError(path)

    loader = _REGISTRY.get(path.suffix.lower())
    if loader is None:
        raise UnsupportedFormatError(f"Format non supporté : {path.suffix}")

    return loader(path)