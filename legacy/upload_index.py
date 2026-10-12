from pathlib import Path
import csv
import json
import re

from docx import Document
from pypdf import PdfReader


CHUNK_SIZE = 400
OVERLAP = 80


def clean_text(text):
    text = text.replace("\r\n", "\n")

    # on garde le texte des liens Markdown mais pas l'URL
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)

    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.MULTILINE)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = re.sub(r"[ \t]+", " ", text)

    return text.strip()


def split_into_chunks(text):
    words = text.split()
    chunks = []

    start = 0

    while start < len(words):
        end = min(start + CHUNK_SIZE, len(words))
        chunk = " ".join(words[start:end])

        if chunk.strip():
            chunks.append(chunk)

        if end == len(words):
            break

        # on recule jusqu'à une frontière de phrase pour ne pas couper en plein milieu
        overlap_start = end - OVERLAP

        while overlap_start > start and not words[overlap_start - 1].endswith((".", ":", "!", "?", ";")):
            overlap_start -= 1

        start = overlap_start if overlap_start > start else end - OVERLAP

    return chunks


def read_txt_or_md(path):
    raw = path.read_text(encoding="utf-8", errors="ignore").strip()

    # un fichier qui commence par { est probablement un JSON du crawler
    if raw.startswith("{"):
        try:
            content = json.loads(raw)
        except json.JSONDecodeError:
            return raw

        if isinstance(content, dict) and isinstance(content.get("text"), str):
            return content["text"]

    return raw


def read_pdf(path):
    reader = PdfReader(str(path))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        pages.append(f"[Page {page_number}]\n{text}")

    return "\n\n".join(pages)


def read_docx(path):
    document = Document(str(path))
    return "\n".join(
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    )


def read_csv_file(path):
    rows = []

    with path.open("r", encoding="utf-8", errors="ignore", newline="") as file:
        reader = csv.reader(file)

        for row in reader:
            rows.append(" | ".join(row))

    return "\n".join(rows)


def read_xlsx(path):
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), read_only=True, data_only=True)
    lines = []

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(cell) for cell in row if cell is not None]

            if cells:
                lines.append(" | ".join(cells))

    workbook.close()
    return "\n".join(lines)


def extract_text(path):
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md", ".json"}:
        return read_txt_or_md(path)

    if suffix == ".pdf":
        return read_pdf(path)

    if suffix == ".docx":
        return read_docx(path)

    if suffix == ".csv":
        return read_csv_file(path)

    if suffix == ".xlsx":
        return read_xlsx(path)

    raise ValueError(f"Format non supporté : {suffix}")


def process_files(files):
    documents = []
    seen_chunks = set()

    for file in files:
        path = Path(file)

        try:
            text = extract_text(path)
        except Exception as error:
            print(f"Erreur avec {path.name} : {error}")
            continue

        chunks = split_into_chunks(clean_text(text))

        for chunk_number, chunk in enumerate(chunks):
            # évite d'indexer deux fois le même passage
            fingerprint = hash(chunk[:200])

            if fingerprint in seen_chunks:
                continue

            seen_chunks.add(fingerprint)

            documents.append({
                "id": f"{path.stem}-{chunk_number}",
                "source": path.name,
                "path": str(path),
                "chunk_id": chunk_number,
                "text": chunk,
            })

    output_dir = Path("data/processed")
    output_dir.mkdir(parents=True, exist_ok=True)

    output_file = output_dir / "uploaded_chunks.json"
    output_file.write_text(
        json.dumps(documents, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    return documents
