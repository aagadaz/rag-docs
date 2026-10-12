import pytest
from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation

from rag.loaders import UnsupportedFormatError, load_source, supported_extensions

SIMPLE_TYPES = (str, int, float, bool)


def _assert_valid(docs):
    assert docs, "le loader doit renvoyer au moins un Document"
    for doc in docs:
        assert doc.page_content.strip()
        assert "source" in doc.metadata
        assert "file_type" in doc.metadata
        for value in doc.metadata.values():
            assert isinstance(value, SIMPLE_TYPES), f"métadonnée non supportée par Chroma : {value!r}"


@pytest.fixture
def sample_files(tmp_path):
    files = {}

    files["txt"] = tmp_path / "notes.txt"
    files["txt"].write_text("Bonjour\n\nDeuxième paragraphe.", encoding="utf-8")

    files["py"] = tmp_path / "script.py"
    files["py"].write_text("def f():\n    return 1\n", encoding="utf-8")

    files["json"] = tmp_path / "data.json"
    files["json"].write_text('{"a": [1, 2]}', encoding="utf-8")

    files["csv"] = tmp_path / "table.csv"
    files["csv"].write_text("Modèle,Tension\nB,12\n", encoding="cp1252")

    files["html"] = tmp_path / "page.html"
    files["html"].write_text(
        "<html><head><title>Doc</title></head><body><nav>menu</nav>"
        "<main><h1>Titre</h1><pre>x = 1\n    y = 2</pre></main></body></html>",
        encoding="utf-8",
    )

    docx_file = DocxDocument()
    docx_file.add_heading("Titre", level=1)
    docx_file.add_paragraph("Un paragraphe.")
    table = docx_file.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    files["docx"] = tmp_path / "doc.docx"
    docx_file.save(files["docx"])

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Slide 1"
    files["pptx"] = tmp_path / "deck.pptx"
    presentation.save(files["pptx"])

    workbook = Workbook()
    workbook.active.append(["Modèle", "Tension"])
    workbook.active.append(["B", 12])
    files["xlsx"] = tmp_path / "sheet.xlsx"
    workbook.save(files["xlsx"])

    return files


@pytest.mark.parametrize("kind", ["txt", "py", "json", "csv", "html", "docx", "pptx", "xlsx"])
def test_every_format_returns_valid_documents(sample_files, kind):
    _assert_valid(load_source(sample_files[kind]))


def test_docx_keeps_headings_and_tables(sample_files):
    text = load_source(sample_files["docx"])[0].page_content
    assert "# Titre" in text
    assert "A | B" in text


def test_code_keeps_indentation(sample_files):
    text = load_source(sample_files["py"])[0].page_content
    assert "    return 1" in text


def test_html_removes_navigation_and_keeps_code(sample_files):
    text = load_source(sample_files["html"])[0].page_content
    assert "menu" not in text
    assert "```" in text
    assert "    y = 2" in text


def test_csv_rows_are_self_describing(sample_files):
    text = load_source(sample_files["csv"])[0].page_content
    assert "Modèle: B" in text


def test_unknown_extension_raises(tmp_path):
    path = tmp_path / "fichier.xyz"
    path.write_text("contenu")
    with pytest.raises(UnsupportedFormatError):
        load_source(path)


def test_missing_file_raises(tmp_path):
    with pytest.raises(FileNotFoundError):
        load_source(tmp_path / "absent.pdf")


def test_supported_extensions_lists_main_formats():
    extensions = supported_extensions()
    for ext in [".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".html", ".md", ".py"]:
        assert ext in extensions